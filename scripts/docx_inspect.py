#!/usr/bin/env python3
"""Strukturierte DOCX-Inspektion (Outline, Element-Details, Substring-Suche).

Inspiriert vom view/get-Konzept aus OfficeCLI, nativ mit python-docx umgesetzt.
Dient /apply-feedback (Anker-Lokalisierung + Run-Struktur) und schnellen
Struktur-Checks ohne Word.

Aufrufe:
  python3 scripts/docx_inspect.py outline <datei>.docx [--json]
  python3 scripts/docx_inspect.py get <datei>.docx <adresse> [--json]
  python3 scripts/docx_inspect.py search <datei>.docx "<substring>" [--json]

Adressen:
  p<N>            Body-Absatz N (0-basiert, wie in outline/search angezeigt)
  t<N>            Tabelle N
  t<N>/r<R>/c<C>  Zelle (Zeile R, Spalte C) der Tabelle N
"""

import argparse
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def _run_info(run):
    return {
        "text": run.text,
        "bold": run.bold,
        "italic": run.italic,
        "underline": run.underline,
        "font": run.font.name,
        "size_pt": run.font.size.pt if run.font.size else None,
    }


def _para_info(idx, para, address):
    pf = para.paragraph_format
    return {
        "address": address,
        "index": idx,
        "style": para.style.name if para.style else None,
        "text": para.text,
        "runs": [_run_info(r) for r in para.runs],
        "line_spacing": pf.line_spacing,
        "has_section_break": para._p.find(qn("w:pPr")) is not None
        and para._p.find(qn("w:pPr")).find(qn("w:sectPr")) is not None,
        "has_page_break": any(
            br.get(qn("w:type")) == "page" for br in para._p.iter(qn("w:br"))
        ),
    }


def _is_toc_field(para):
    for instr in para._p.iter(qn("w:instrText")):
        if instr.text and "TOC" in instr.text:
            return True
    return False


def iter_addressed_paragraphs(doc):
    """Alle Absätze mit Adresse: Body-Absätze und Tabellenzellen-Absätze."""
    for i, p in enumerate(doc.paragraphs):
        yield f"p{i}", p
    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                for pi, p in enumerate(cell.paragraphs):
                    yield f"t{ti}/r{ri}/c{ci}/p{pi}", p


def cmd_outline(doc, as_json):
    items = []
    heading_re = re.compile(r"^Heading (\d)$")
    for i, p in enumerate(doc.paragraphs):
        style = p.style.name if p.style else ""
        m = heading_re.match(style)
        entry = {
            "address": f"p{i}",
            "style": style,
            "level": int(m.group(1)) if m else None,
            "text": p.text.strip(),
            "runs": len(p.runs),
            "toc_field": _is_toc_field(p),
        }
        if entry["text"] or entry["toc_field"]:
            items.append(entry)
    tables = [
        {"address": f"t{ti}", "rows": len(t.rows), "cols": len(t.columns)}
        for ti, t in enumerate(doc.tables)
    ]
    if as_json:
        print(json.dumps({"paragraphs": items, "tables": tables},
                         ensure_ascii=False, indent=2))
        return
    for e in items:
        if e["toc_field"]:
            print(f"{e['address']:>6}  [TOC-Feld]")
            continue
        indent = "  " * (e["level"] or 0) if e["level"] else "      "
        marker = f"H{e['level']} " if e["level"] else ""
        text = e["text"][:90]
        print(f"{e['address']:>6}  {indent}{marker}{text}  "
              f"({e['style']}, {e['runs']} Runs)")
    for t in tables:
        print(f"{t['address']:>6}  [Tabelle {t['rows']}×{t['cols']}]")


def _resolve_address(doc, address):
    m = re.fullmatch(r"p(\d+)", address)
    if m:
        idx = int(m.group(1))
        if idx >= len(doc.paragraphs):
            sys.exit(f"FEHLER: p{idx} existiert nicht "
                     f"(nur {len(doc.paragraphs)} Body-Absätze).")
        return ("para", idx, doc.paragraphs[idx])
    m = re.fullmatch(r"t(\d+)", address)
    if m:
        idx = int(m.group(1))
        if idx >= len(doc.tables):
            sys.exit(f"FEHLER: t{idx} existiert nicht "
                     f"(nur {len(doc.tables)} Tabellen).")
        return ("table", idx, doc.tables[idx])
    m = re.fullmatch(r"t(\d+)/r(\d+)/c(\d+)(?:/p(\d+))?", address)
    if m:
        ti, ri, ci = int(m.group(1)), int(m.group(2)), int(m.group(3))
        try:
            cell = doc.tables[ti].rows[ri].cells[ci]
        except IndexError:
            sys.exit(f"FEHLER: Adresse {address} existiert nicht.")
        if m.group(4) is not None:
            pi = int(m.group(4))
            if pi >= len(cell.paragraphs):
                sys.exit(f"FEHLER: {address} existiert nicht.")
            return ("para", address, cell.paragraphs[pi])
        return ("cell", address, cell)
    sys.exit(f"FEHLER: Unbekanntes Adressformat: {address!r} "
             "(erwartet p<N>, t<N> oder t<N>/r<R>/c<C>[/p<P>]).")


def cmd_get(doc, address, as_json):
    kind, idx, obj = _resolve_address(doc, address)
    if kind == "para":
        info = _para_info(idx, obj, address)
    elif kind == "cell":
        info = {
            "address": address,
            "paragraphs": [
                _para_info(pi, p, f"{address}/p{pi}")
                for pi, p in enumerate(obj.paragraphs)
            ],
        }
    else:
        info = {
            "address": address,
            "rows": len(obj.rows),
            "cols": len(obj.columns),
            "cells": [
                [cell.text for cell in row.cells] for row in obj.rows
            ],
            "style": obj.style.name if obj.style else None,
        }
    if as_json:
        print(json.dumps(info, ensure_ascii=False, indent=2))
    else:
        print(json.dumps(info, ensure_ascii=False, indent=2))


def cmd_search(doc, needle, as_json):
    hits = []
    for address, p in iter_addressed_paragraphs(doc):
        if needle not in p.text:
            continue
        run_concat = "".join(r.text for r in p.runs)
        hits.append({
            "address": address,
            "style": p.style.name if p.style else None,
            "text": p.text,
            "run_count": len(p.runs),
            "needle_in_single_run": any(needle in r.text for r in p.runs),
            "needle_spans_runs": needle in run_concat
            and not any(needle in r.text for r in p.runs),
            "needle_outside_runs": needle not in run_concat,
            "run_texts": [r.text for r in p.runs],
        })
    if as_json:
        print(json.dumps({"needle": needle, "hits": hits},
                         ensure_ascii=False, indent=2))
        return
    if not hits:
        print(f"Keine Treffer für {needle!r}.")
        sys.exit(1)
    print(f"{len(hits)} Treffer für {needle!r}:\n")
    for h in hits:
        flags = []
        if h["needle_in_single_run"]:
            flags.append("1-Run-Patch möglich")
        if h["needle_spans_runs"]:
            flags.append("über Run-Grenzen (Multi-Run-Replace nötig)")
        if h["needle_outside_runs"]:
            flags.append("außerhalb normaler Runs (z. B. Hyperlink)")
        print(f"  {h['address']}  ({h['style']}, {h['run_count']} Runs"
              f"{'; ' + ', '.join(flags) if flags else ''})")
        print(f"    {h['text'][:160]}")


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("command", choices=["outline", "get", "search"])
    ap.add_argument("docx", type=Path)
    ap.add_argument("arg", nargs="?",
                    help="Adresse (get) bzw. Substring (search)")
    ap.add_argument("--json", action="store_true", dest="as_json")
    args = ap.parse_args()

    if not args.docx.exists():
        sys.exit(f"FEHLER: Datei nicht gefunden: {args.docx}")
    doc = Document(str(args.docx))

    if args.command == "outline":
        cmd_outline(doc, args.as_json)
    elif args.command == "get":
        if not args.arg:
            sys.exit("FEHLER: get braucht eine Adresse (z. B. p42).")
        cmd_get(doc, args.arg, args.as_json)
    else:
        if not args.arg:
            sys.exit("FEHLER: search braucht einen Substring.")
        cmd_search(doc, args.arg, args.as_json)


if __name__ == "__main__":
    main()
