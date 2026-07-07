#!/usr/bin/env python3
"""Kanonischer IU-DOCX-Builder (config-parametrisiert).

Baut aus den fertigen Kapiteln (output/phase-05-writing/final/*.md) ein
IU-formatkonformes DOCX: DIN A4, 2 cm Ränder, Arial 11/1,5, schwarze
Überschriften, römische/arabische Seitenzählung, TOC-/Tabellen-/Abkürzungs-/
Literaturverzeichnis, hängender Einzug im Lit-Verz.

Alle projektspezifischen Werte kommen aus config.yaml — keine hartcodierten
Namen/Kurse/Pfade mehr (Vorgänger: fallstudie-blockchain/scripts/build_iu_docx.py).

Aufruf:
    python3 scripts/build_docx.py [PROJEKT_DIR]
    (ohne Argument: aktuelles Arbeitsverzeichnis)

Optionale Projekt-Dateien:
    output/terminology.md   # Abkürzungen (## Abkürzungen, dann "- ABK — Bedeutung")
    output/tables.yaml       # {"3.4": "Caption-Text"} → Tabellen-Beschriftungen je Abschnitt
"""
import re
import sys
import subprocess
import shutil
import yaml
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ========== Projekt + Config ==========

PROJECT = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else Path.cwd()
config = yaml.safe_load((PROJECT / "config.yaml").read_text()) or {}


def cfg(path, default=None):
    """Liest verschachtelten config-Wert per Punkt-Pfad, mit Fallback."""
    cur = config
    for key in path.split("."):
        if isinstance(cur, dict) and cur.get(key) is not None:
            cur = cur[key]
        else:
            return default
    return cur


# Titelblatt-Metadaten
autor_name = cfg("autor.name", "")
matrikel = cfg("autor.matrikelnummer", "")
hochschule = cfg("autor.hochschule", "IU Internationale Hochschule")
studiengang = cfg("autor.studiengang", "")
erstgut = cfg("betreuung.erstgutachter", "")
titel = cfg("projekt.titel", "")
abgabedat = cfg("abgabe.datum") or "[Abgabedatum ergänzen]"
kurs_code = cfg("projekt.kurs_modul", "")
kurs_name = cfg("projekt.kurs_titel", "")
semester = cfg("projekt.semester", "")
typ = (cfg("projekt.typ", "fallstudie") or "fallstudie")
pruefungsform = "Fallstudie" if typ == "fallstudie" else typ.capitalize()

# Formatierung (IU-Defaults)
FONT = cfg("formatierung.schriftart", "Arial")
FSIZE = int(cfg("formatierung.schriftgroesse", 11))
TFSIZE = int(cfg("formatierung.tabellen_schriftgroesse", 10))
LINESP = float(cfg("formatierung.zeilenabstand", 1.5))
M_TOP = float(cfg("formatierung.seitenränder.oben", 2.0))
M_BOT = float(cfg("formatierung.seitenränder.unten", 2.0))
M_LEFT = float(cfg("formatierung.seitenränder.links", 2.0))
M_RIGHT = float(cfg("formatierung.seitenränder.rechts", 2.0))

# Verzeichnisse
SHOW_TOC = cfg("verzeichnisse.inhaltsverzeichnis", True)
SHOW_TAB = cfg("verzeichnisse.tabellenverzeichnis", False)
SHOW_ABK = cfg("verzeichnisse.abkuerzungsverzeichnis", True)
SHOW_LIT = cfg("verzeichnisse.literaturverzeichnis", True)

# Logo (config → Auto-Detect → keins)
logo_rel = cfg("logo.pfad", "")
LOGO = (PROJECT / logo_rel) if logo_rel else None
if not (LOGO and LOGO.exists()):
    LOGO = None
    for cand in ("iu-logo.png", "logo.png", "logo.jpg", "logo.jpeg"):
        p = PROJECT / "assets" / "img" / cand
        if p.exists():
            LOGO = p
            break

# Tabellen-Captions (optional)
TABLE_CAPTIONS = {}
_tables_file = PROJECT / "output" / "tables.yaml"
if _tables_file.exists():
    TABLE_CAPTIONS = yaml.safe_load(_tables_file.read_text()) or {}

CHAP_DIR = PROJECT / "output" / "phase-05-writing" / "final"
OUTPUT = PROJECT / "output" / "phase-07-docx"
OUTPUT.mkdir(parents=True, exist_ok=True)

# ========== Helper ==========

def set_font(run, name=None, size=None, bold=False, italic=False):
    name = name or FONT
    size = size or FSIZE
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), name)


def add_para(doc, text, bold=False, italic=False, size=None, align=None,
             space_before=0, space_after=0, line_spacing=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing = line_spacing or LINESP
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    for i, line in enumerate(text.split("\n")):
        if i > 0:
            p.add_run().add_break()
        set_font(p.add_run(line), FONT, size, bold, italic)
    return p


def add_empty_para(doc, count=1, size=None):
    for _ in range(count):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        set_font(p.add_run(), FONT, size)


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def set_margins(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(M_TOP)
    section.bottom_margin = Cm(M_BOT)
    section.left_margin = Cm(M_LEFT)
    section.right_margin = Cm(M_RIGHT)


def set_page_number_format(section, fmt="decimal", start=None):
    sectPr = section._sectPr
    pgNumType = sectPr.find(qn("w:pgNumType"))
    if pgNumType is None:
        pgNumType = OxmlElement("w:pgNumType")
        sectPr.append(pgNumType)
    pgNumType.set(qn("w:fmt"), fmt)
    if start is not None:
        pgNumType.set(qn("w:start"), str(start))


def add_page_number_footer(section):
    footer = section.footer
    for para in footer.paragraphs:
        para.clear()
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_font(run, FONT, FSIZE)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._element.append(fld1)
    run._element.append(instr)
    run._element.append(fld2)


def add_field(run, instr_text, placeholder):
    """Word-Feld (TOC/SEQ) mit Platzhalter einfügen."""
    fc1 = OxmlElement("w:fldChar"); fc1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr_text
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    ph = OxmlElement("w:t"); ph.text = placeholder
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for el in (fc1, it, sep, ph, end):
        run._element.append(el)


# ========== Dokument-Grundgerüst ==========

doc = Document()
for section in doc.sections:
    set_margins(section)

for h_name, h_size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 11)):
    style = doc.styles[h_name]
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.name = FONT
    style.font.size = Pt(h_size)
    style.font.bold = True

from docx.enum.style import WD_STYLE_TYPE
# Inhaltsverzeichnis-Ebenen: Ebene 1 fett (IU-Vorgabe), Ebenen 2 und 3 normal.
# TOC 2/3 explizit anlegen, damit Word sie beim Feld-Update nicht inkonsistent
# selbst erzeugt und der Kontrast fett/normal garantiert ist.
for _toc_name, _toc_bold in (("TOC 1", True), ("TOC 2", False), ("TOC 3", False)):
    try:
        _toc_style = doc.styles[_toc_name]
    except KeyError:
        _toc_style = doc.styles.add_style(_toc_name, WD_STYLE_TYPE.PARAGRAPH)
    _toc_style.font.bold = _toc_bold
    _toc_style.font.name = FONT
    _toc_style.font.size = Pt(FSIZE)
    _toc_style.font.color.rgb = RGBColor(0, 0, 0)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(FSIZE)
_rPr = normal.element.get_or_add_rPr()
_rFonts = _rPr.find(qn("w:rFonts"))
if _rFonts is None:
    _rFonts = OxmlElement("w:rFonts")
    _rPr.append(_rFonts)
for _a in ("w:ascii", "w:hAnsi", "w:cs"):
    _rFonts.set(qn(_a), FONT)
normal.paragraph_format.line_spacing = LINESP
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(0)

# Automatische Silbentrennung: ohne autoHyphenation reisst Word/LibreOffice im
# Blocksatz grosse Wortzwischenraeume. Im settings-Part setzen, damit der ganze
# Body getrennt wird; Trennzone 0,5 cm, max. zwei Trennstriche in Folge.
_settings = doc.settings.element
for _tag, _attrs in (
    ("w:autoHyphenation", {"w:val": "true"}),
    ("w:consecutiveHyphenLimit", {"w:val": "2"}),
    ("w:hyphenationZone", {"w:val": "284"}),  # 0,5 cm in Twips
):
    if _settings.find(qn(_tag)) is None:
        _el = OxmlElement(_tag)
        for _k, _v in _attrs.items():
            _el.set(qn(_k), _v)
        _settings.append(_el)

# ========== Titelblatt ==========

section0 = doc.sections[0]
section0.different_first_page_header_footer = True
add_empty_para(doc, 1)

if LOGO:
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.add_run().add_picture(str(LOGO), width=Cm(10))

add_para(doc, hochschule, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
add_empty_para(doc, 3)
add_para(doc, f"Prüfungsform: {pruefungsform}", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
add_empty_para(doc, 1)
add_para(doc, titel, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
add_empty_para(doc, 1)
add_para(doc, f"Verfasser: {autor_name}", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
add_empty_para(doc, 2)

modul_zeile = f"Modul: {kurs_code}" + (f" – {kurs_name}" if kurs_name else "")
for zeile in (
    modul_zeile,
    f"Tutor: {erstgut}" if erstgut else None,
    f"Semester: {semester}" if semester else None,
    f"Studiengang: {studiengang}" if studiengang else None,
    f"Matrikelnummer: {matrikel}",
    f"Eingereicht am: {abgabedat}",
):
    if zeile:
        add_para(doc, zeile, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)

# ========== Frontmatter (römisch) ==========

fm = doc.add_section(WD_SECTION.NEW_PAGE)
set_margins(fm)
set_page_number_format(fm, fmt="upperRoman", start=2)
# Erste Seite NICHT als Sonderseite behandeln: erbt sonst die unterdrueckte
# Titelblatt-Fusszeile, sodass das Inhaltsverzeichnis keine "II" zeigt.
fm.different_first_page_header_footer = False
fm.footer.is_linked_to_previous = False
add_page_number_footer(fm)
section0.footer.is_linked_to_previous = False

if SHOW_TOC:
    p_toc_title = doc.add_paragraph()
    p_toc_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    # Bewusst KEIN Heading 1: sonst listet sich das Inhaltsverzeichnis selbst.
    # Optik (16pt, fett, schwarz, linksbuendig) gleicht aber Heading 1 an.
    r = p_toc_title.add_run("Inhaltsverzeichnis")
    set_font(r, FONT, 16, bold=True); r.font.color.rgb = RGBColor(0, 0, 0)
    p_toc_title.paragraph_format.space_after = Pt(18)
    p_toc_field = doc.add_paragraph()
    p_toc_field.style = doc.styles["TOC 1"]
    r = p_toc_field.add_run(); set_font(r, FONT, FSIZE, bold=True)
    add_field(r, r'TOC \o "1-3" \h \z \u',
              "Rechtsklick → Feld aktualisieren, um das Inhaltsverzeichnis anzuzeigen.")

if SHOW_TAB:
    add_empty_para(doc, 1)
    p_tab_title = doc.add_paragraph()
    # Heading 1, damit das Tabellenverzeichnis als Eintrag im Inhaltsverzeichnis erscheint.
    p_tab_title.style = doc.styles["Heading 1"]
    rr = p_tab_title.add_run("Tabellenverzeichnis")
    set_font(rr, FONT, 14, bold=True); rr.font.color.rgb = RGBColor(0, 0, 0)
    p_tab_title.paragraph_format.space_before = Pt(18)
    p_tab_title.paragraph_format.space_after = Pt(12)
    r = doc.add_paragraph().add_run(); set_font(r, FONT, FSIZE)
    add_field(r, r'TOC \h \z \c "Tabelle"',
              "Rechtsklick → Feld aktualisieren, um das Tabellenverzeichnis anzuzeigen.")

term_file = PROJECT / "output" / "terminology.md"
if SHOW_ABK and term_file.exists():
    add_page_break(doc)
    p_abk_title = doc.add_paragraph()
    p_abk_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_abk_title.style = doc.styles["Heading 1"]
    r = p_abk_title.add_run("Abkürzungsverzeichnis")
    set_font(r, FONT, 16, bold=True); r.font.color.rgb = RGBColor(0, 0, 0)
    p_abk_title.paragraph_format.space_after = Pt(18)
    term_content = term_file.read_text()
    abk_match = re.search(r"## Abk.rzungen.*?\n(.+?)(?=\n##|\Z)", term_content, re.DOTALL)
    abk_lines = re.findall(r"^- (\S+)\s+[—-]+\s+(.+)$", abk_match.group(1), re.MULTILINE) if abk_match else []
    if abk_lines:
        table = doc.add_table(rows=len(abk_lines), cols=2)
        table.autofit = False
        for row in table.rows:
            row.cells[0].width = Cm(4)
            row.cells[1].width = Cm(13)
        for i, (abk, bedeut) in enumerate(abk_lines):
            c0 = table.rows[i].cells[0]; c0.text = ""
            set_font(c0.paragraphs[0].add_run(abk), FONT, FSIZE, bold=True)
            c1 = table.rows[i].cells[1]; c1.text = ""
            set_font(c1.paragraphs[0].add_run(bedeut), FONT, FSIZE)

# ========== Body (arabisch) ==========

body = doc.add_section(WD_SECTION.NEW_PAGE)
set_margins(body)
set_page_number_format(body, fmt="decimal", start=1)
# Erste Body-Seite (Einleitung) soll "1" tragen, daher keine Sonderbehandlung
# der ersten Seite, sonst fehlt die Seitenzahl auf Seite 1.
body.different_first_page_header_footer = False
body.footer.is_linked_to_previous = False
add_page_number_footer(body)


def add_table_caption(doc, table_index, caption_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing = 1.15; pf.space_before = Pt(2); pf.space_after = Pt(8)
    try:
        p.style = doc.styles["Caption"]
    except KeyError:
        pass
    r1 = p.add_run("Tabelle "); set_font(r1, FONT, 10, bold=True); r1.font.color.rgb = RGBColor(0, 0, 0)
    r_seq = p.add_run(); set_font(r_seq, FONT, 10, bold=True); r_seq.font.color.rgb = RGBColor(0, 0, 0)
    add_field(r_seq, r" SEQ Tabelle \* ARABIC ", str(table_index))
    r2 = p.add_run(f": {caption_text}"); set_font(r2, FONT, 10); r2.font.color.rgb = RGBColor(0, 0, 0)


TWIPS_PER_CM = 566.93


def compute_col_widths_cm(header_cells, body_rows, total_cm, min_cm=1.0, cap=100):
    """Inhaltsproportionale Spaltenbreiten (Summe = total_cm).

    Jede Spalte erhaelt mindestens min_cm; der Rest wird proportional zur
    laengsten Zelle je Spalte verteilt (gedeckelt bei cap Zeichen, damit eine
    sehr lange Textspalte die anderen nicht ausstarrt). Verhindert, dass eine
    lange Spalte (z. B. die User-Story-Spalte) bei autofit gleich schmal wie
    eine winzige Spalte wird.
    """
    n = len(header_cells)
    if n == 0:
        return []
    nat = []
    for j in range(n):
        cells = [header_cells[j]] + [r[j] for r in body_rows if j < len(r)]
        nat.append(max((len(c.strip()) for c in cells), default=1))
    capped = [min(c, cap) for c in nat]
    s = sum(capped) or 1
    remaining = max(0.0, total_cm - min_cm * n)
    return [min_cm + remaining * (c / s) for c in capped]


def add_markdown_table(doc, header_cells, body_rows, caption=None, table_index=1):
    n_cols = len(header_cells)
    table = doc.add_table(rows=1 + len(body_rows), cols=n_cols)
    # autofit=False -> <w:tblLayout w:type="fixed"/> (an schema-korrekter Stelle
    # durch python-docx eingefuegt). Nur mit fixed layout greifen feste Breiten.
    table.autofit = False
    text_width_cm = 21.0 - M_LEFT - M_RIGHT
    widths = compute_col_widths_cm(header_cells, body_rows, text_width_cm)

    def _fill(cell, txt, bold=False):
        cell.text = ""
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.space_before = Pt(0); pf.space_after = Pt(0)
        run = p.add_run(txt.strip()); set_font(run, FONT, TFSIZE, bold=bold)
        run.font.color.rgb = RGBColor(0, 0, 0)

    for j, txt in enumerate(header_cells):
        _fill(table.rows[0].cells[j], txt, bold=True)
    for i, row in enumerate(body_rows, start=1):
        for j, txt in enumerate(row):
            if j < n_cols:
                _fill(table.rows[i].cells[j], txt)

    tbl = table._tbl
    tblPr = tbl.find(qn("w:tblPr"))
    # Gesamtbreite der Tabelle fixieren (dxa)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.insert_element_before(tblW, "w:tblLayout", "w:tblCellMar", "w:tblLook")
    tblW.set(qn("w:w"), str(int(round(sum(widths) * TWIPS_PER_CM))))
    tblW.set(qn("w:type"), "dxa")
    # Spaltenbreiten im tblGrid setzen ...
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for gc, wcm in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(round(wcm * TWIPS_PER_CM))))
    # ... und zusaetzlich je Zelle (tcW), damit Word und LibreOffice folgen.
    all_rows = table.rows
    for row in all_rows:
        for cell, wcm in zip(row.cells, widths):
            cell.width = Cm(wcm)
    # Tabelle nicht ueber den Seitenumbruch zerreissen: jede Zeile cantSplit
    # (kein Bruch mitten in einer Zeile) und alle Zeilen ausser der letzten
    # "mit naechster zusammenhalten" -> Word schiebt die ganze Tabelle bei
    # Bedarf geschlossen auf die naechste Seite, statt sie zu teilen.
    for ri, row in enumerate(all_rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if ri < len(all_rows) - 1:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True
    # Rahmen: oben/unten + Kopf-Trennlinie kraeftiger (sz 8), dazu duenne Linien
    # zwischen Zeilen UND Spalten (insideH/insideV, sz 4); aeussere Seiten offen.
    tblBorders = OxmlElement("w:tblBorders")
    for tag, val, sz in (
        ("top", "single", "8"), ("left", "nil", None),
        ("bottom", "single", "8"), ("right", "nil", None),
        ("insideH", "single", "4"), ("insideV", "single", "4"),
    ):
        b = OxmlElement(f"w:{tag}"); b.set(qn("w:val"), val)
        if sz:
            b.set(qn("w:sz"), sz); b.set(qn("w:color"), "000000")
        tblBorders.append(b)
    tblPr.insert_element_before(
        tblBorders, "w:shd", "w:tblLayout", "w:tblCellMar",
        "w:tblLook", "w:tblCaption", "w:tblDescription",
    )
    if caption:
        add_table_caption(doc, table_index, caption)


def process_chapter_line(doc, line):
    if not line.strip():
        return
    h_match = re.match(r"^(#+)\s+(.+)$", line)
    if h_match:
        level = len(h_match.group(1))
        text = h_match.group(2).strip()
        adj_level = max(1, level - 1)
        size = {1: 16, 2: 14, 3: 11}.get(adj_level, 11)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.space_before = Pt(12); pf.space_after = Pt(12 if adj_level == 1 else 6); pf.line_spacing = 1.15
        p.add_run(text)
        p.style = doc.styles[f"Heading {min(adj_level, 3)}"]
        for r in p.runs:
            set_font(r, FONT, size, bold=True); r.font.color.rgb = RGBColor(0, 0, 0)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing = LINESP; pf.space_before = Pt(0); pf.space_after = Pt(6)
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", line):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            set_font(p.add_run(part[2:-2]), FONT, FSIZE, bold=True)
        elif part.startswith("*") and part.endswith("*"):
            set_font(p.add_run(part[1:-1]), FONT, FSIZE, italic=True)
        else:
            set_font(p.add_run(part), FONT, FSIZE)


all_chapter_text = ""
# Strikte Whitelist statt "*.md minus .prehum.": nur echte Kapiteldateien
# (kapitel-<Zahl>.md, optional kapitel-1.2.md) bauen. Sonst geraten Backups wie
# kapitel-2.precodex.md, kapitel-3.precodex2.md, kapitel-4.md.bak3 oder
# .prehum.-Stände als zusätzliche „Kapitel" in den Build.
_chapter_re = re.compile(r"^kapitel-\d+(?:[.\-]\d+)*\.md$")
for cf in sorted(p for p in CHAP_DIR.glob("kapitel-*.md") if _chapter_re.match(p.name)):
    content = re.sub(r"^---\n.*?\n---\n", "", cf.read_text(), count=1, flags=re.DOTALL)
    all_chapter_text += content + "\n\n"

table_counter = 0
last_section_id = None
for para in re.split(r"\n\s*\n", all_chapter_text):
    para = para.strip()
    if not para:
        continue
    sec_match = re.match(r"^###\s+(\d+(?:\.\d+)?)", para)
    if sec_match:
        last_section_id = sec_match.group(1)
    if para.startswith("|") and re.search(r"^\|[\s:|-]+\|\s*$", para, flags=re.MULTILINE):
        rows = [r for r in para.split("\n") if r.strip().startswith("|")]
        rows = [r for r in rows if not re.match(r"^\|[\s:|-]+\|\s*$", r)]
        parsed = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
        if len(parsed) >= 2:
            table_counter += 1
            add_markdown_table(doc, parsed[0], parsed[1:],
                               caption=TABLE_CAPTIONS.get(last_section_id), table_index=table_counter)
            continue
    lines = para.split("\n")
    if len(lines) == 1:
        process_chapter_line(doc, para)
    else:
        for h in (l for l in lines if re.match(r"^#+\s", l)):
            process_chapter_line(doc, h)
        body_lines = [l for l in lines if not re.match(r"^#+\s", l) and l.strip()]
        if body_lines:
            process_chapter_line(doc, " ".join(body_lines))

# ========== Literaturverzeichnis ==========

def parse_literature():
    lit_file = PROJECT / "sources" / "literature.md"
    if not lit_file.exists():
        return []
    content = lit_file.read_text()
    parts = re.split(r"#\s*PART 2", content, maxsplit=1)
    blocks = re.findall(r"^---\n(.+?)\n---", parts[0], re.MULTILINE | re.DOTALL)
    entries = []
    for block in blocks:
        try:
            data = yaml.safe_load(block)
            if isinstance(data, dict) and "quelle_id" in data:
                entries.append(data)
        except Exception:
            pass
    return entries


def format_authors_apa(autor_str):
    authors = [a.strip() for a in autor_str.split("/")]
    if len(authors) == 1:
        return authors[0]
    if len(authors) == 2:
        return f"{authors[0]}, & {authors[1]}"
    return ", ".join(authors[:-1]) + f", & {authors[-1]}"


def get_year_with_suffix(e):
    m = re.search(r"_(\d{4})([a-z])$", e.get("quelle_id", ""))
    return f"{e.get('jahr', '')}{m.group(2) if m else ''}"


def is_institution(autor_full):
    return "," not in autor_full.split("/")[0]


def _cite_in_text(patterns, text):
    """Prüft case-insensitiv, ob eines der In-Text-Zitatmuster im Volltext steht.

    Ein satzinitial großgeschriebener Partikel-Nachname („De Vries (2018") muss
    dieselbe Quelle treffen wie die Stammform aus literature.md („de Vries"),
    sonst fällt eine zitierte Quelle aus dem Literaturverzeichnis (echter
    APA-Fehler). Die negative Lookbehind ``(?<!\\w)`` verhindert, dass ein kurzes
    Kürzel mitten in einem Wort matcht (z. B. „un, 2015" in „Jun, 2015" für eine
    UN-Quelle); die Muster bleiben durch Jahr + „(" bzw. „, " verankert, daher
    sind Falsch-Positive nahe null.
    """
    for p in patterns:
        if re.search(r"(?<!\w)" + re.escape(p), text, re.IGNORECASE):
            return True
    return False


def get_used_sources(chapters_text):
    """Generisch (datengetrieben, keine hartcodierten Institutionsnamen)."""
    used = []
    for e in parse_literature():
        autor_full = e.get("autor", "")
        year = get_year_with_suffix(e)
        patterns = []
        if is_institution(autor_full):
            name = autor_full.split("/")[0].strip()
            patterns += [f"{name} ({year}", f"{name}, {year}"]
            # Abkürzung: explizites abk-Feld bevorzugen, sonst dynamisch aus dem Text
            # lernen ("Europäische Zentralbank [EZB] ..." → EZB). Kein hartcodiertes Mapping.
            abk = e.get("abk", "")
            if not abk:
                m = re.search(re.escape(name) + r"\s*\[([A-Za-zÄÖÜ.&/ ]{2,40}?)\]", chapters_text)
                if m:
                    abk = m.group(1).strip()
            if abk:
                patterns += [f"{abk} ({year}", f"{abk}, {year}",
                             f"[{abk}] ({year}", f"[{abk}], {year}"]
        else:
            first = autor_full.split("/")[0].split(",")[0].strip()
            patterns += [f"{first} ({year}", f"{first}, {year}",
                         f"{first} et al. ({year}", f"{first} et al., {year}"]
            if "/" in autor_full:
                surnames = [a.split(",")[0].strip() for a in autor_full.split("/")]
                if len(surnames) == 2:
                    patterns += [f"{surnames[0]} & {surnames[1]} ({year}",
                                 f"{surnames[0]} & {surnames[1]}, {year}",
                                 f"{surnames[0]} und {surnames[1]} ({year}",
                                 f"{surnames[0]} und {surnames[1]}, {year}",
                                 f"{surnames[0]}/{surnames[1]} ({year}"]
                else:
                    patterns.append(f"{first} et al.")
        if _cite_in_text(patterns, chapters_text):
            used.append(e)
    used.sort(key=lambda e: (e.get("autor", "").split("/")[0].split(",")[0].strip().lower(),
                             e.get("quelle_id", "")))
    return used


def title_sep(titel):
    """Trenner nach dem Titel: Punkt + Leerzeichen, ausser der Titel endet bereits
    auf Satzzeichen (?, !, .) — dann nur Leerzeichen, sonst entsteht 'Paradigm?.'."""
    return " " if titel.rstrip().endswith((".", "?", "!")) else ". "


def fmt_pages(seiten):
    """Seitenbereich mit Halbgeviertstrich (en dash) statt Bindestrich:
    'S. 1271-1319' -> 'S. 1271–1319'. Nur den numerischen Bereich umstellen."""
    return re.sub(r"(\d)\s*-\s*(\d)", r"\1–\2", str(seiten))


def add_bib_entry(doc, e):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing = LINESP; pf.space_before = Pt(0); pf.space_after = Pt(10)
    pf.left_indent = Cm(1.27); pf.first_line_indent = Cm(-1.27)

    def r_text(txt, italic=False, bold=False):
        r = p.add_run(txt); set_font(r, FONT, FSIZE, bold=bold, italic=italic)
        r.font.color.rgb = RGBColor(0, 0, 0)

    r_text(f"{format_authors_apa(e.get('autor', ''))} ({get_year_with_suffix(e)}). ")
    titel_e = e.get("titel", "")
    typ_e = e.get("quelle_typ", "")
    if typ_e == "journal":
        r_text(titel_e + title_sep(titel_e))
        r_text(e.get("journal", ""), italic=True)
        if e.get("jahrgang"):
            r_text(", "); r_text(str(e["jahrgang"]), italic=True)
        if e.get("ausgabe"):
            r_text(f"({e['ausgabe']})")
        if e.get("seiten"):
            r_text(f", {fmt_pages(e['seiten'])}")
        r_text(".")
        if e.get("doi"):
            r_text(f" https://doi.org/{e['doi']}")
        elif e.get("url"):
            # Law-/SSRN-Journals ohne DOI (z. B. Arner 2016) sonst ohne stabilen Link.
            r_text(f" {e['url']}")
    elif typ_e == "sammelband":
        r_text(titel_e + title_sep(titel_e))
        r_text("In ")
        r_text(e.get("sammelbandtitel", ""), italic=True)
        if e.get("seiten"):
            r_text(f" (S. {fmt_pages(e['seiten'])})")
        r_text(". ")
        if e.get("verlag"):
            r_text(f"{e['verlag']}.")
        if e.get("doi"):
            r_text(f" https://doi.org/{e['doi']}")
    elif typ_e == "konferenz":
        r_text(titel_e + title_sep(titel_e))
        r_text(e.get("sammelbandtitel", ""), italic=True)
        r_text(". ")
        if e.get("abruf"):
            r_text(f"Abgerufen am {e['abruf']}, von ")
        if e.get("url"):
            r_text(e["url"])
    elif typ_e == "buch":
        r_text(titel_e, italic=True)
        if e.get("auflage"):
            r_text(f" ({e['auflage']}. Aufl.)")
            r_text(". ")
        else:
            r_text(title_sep(titel_e))
        if e.get("ort"):
            r_text(f"{e['ort']}: ")
        r_text(f"{e.get('verlag', '')}.")
    elif typ_e == "website":
        r_text(titel_e, italic=True); r_text(title_sep(titel_e))
        if e.get("verlag"):
            r_text("In "); r_text(e["verlag"], italic=True); r_text(". ")
        if e.get("abruf"):
            r_text(f"Abgerufen am {e['abruf']}, von ")
        if e.get("url"):
            r_text(e["url"])
    elif typ_e in ("report", "working_paper"):
        r_text(titel_e, italic=True); r_text(title_sep(titel_e))
        # Working Paper: Reihe + Nummer gehören VOR Verlag/URL (APA), z. B.
        # „CESifo Working Paper No. 8655." Bisher musste die Reihe ins verlag-Feld
        # geschmuggelt werden — jetzt eigene Felder reihe/nummer.
        reihe = str(e.get("reihe", "")).strip()
        nummer = str(e.get("nummer", "")).strip()
        if reihe or nummer:
            r_text(f"{' '.join(s for s in (reihe, nummer) if s)}. ")
        if e.get("verlag"):
            r_text(f"{e['verlag']}. ")
        if e.get("url"):
            r_text(e["url"])
        elif e.get("doi"):
            r_text(f"https://doi.org/{e['doi']}")
    else:
        r_text(titel_e, italic=True)
        r_text("" if titel_e.rstrip().endswith((".", "?", "!")) else ".")


if SHOW_LIT:
    add_page_break(doc)
    p_lit = doc.add_paragraph()
    p_lit.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_lit.style = doc.styles["Heading 1"]
    p_lit.add_run("Literaturverzeichnis")
    p_lit.paragraph_format.space_before = Pt(12)
    p_lit.paragraph_format.space_after = Pt(18)
    for r in p_lit.runs:
        set_font(r, FONT, 16, bold=True); r.font.color.rgb = RGBColor(0, 0, 0)
    for e in get_used_sources(all_chapter_text):
        add_bib_entry(doc, e)

# ========== Speichern + PDF ==========

def slugify(text):
    s = re.sub(r"[^a-z0-9]+", "-", (text or "").lower()).strip("-")
    return s or "arbeit"


out_name = slugify(kurs_code or titel or "fallstudie")
docx_path = OUTPUT / f"{out_name}.docx"
doc.save(docx_path)
print(f"✓ DOCX gebaut: {docx_path}  ({docx_path.stat().st_size} bytes)")


def export_pdf(docx_path):
    """DOCX → PDF via LibreOffice headless (layout-identisch). Best-effort."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
        soffice = str(mac) if mac.exists() else None
    if not soffice:
        print("⚠ PDF übersprungen: LibreOffice (soffice) nicht gefunden. "
              "DOCX ist das Primärartefakt; PDF später via 'Speichern als PDF'.")
        return None
    try:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf",
                        "--outdir", str(docx_path.parent), str(docx_path)],
                       check=True, capture_output=True, timeout=180)
        pdf_path = docx_path.with_suffix(".pdf")
        if pdf_path.exists():
            print(f"✓ PDF gebaut: {pdf_path}  ({pdf_path.stat().st_size} bytes)")
            return pdf_path
    except Exception as exc:
        print(f"⚠ PDF-Konvertierung fehlgeschlagen: {exc}")
    return None


export_pdf(docx_path)
print()
print("Nächste Schritte vor Abgabe:")
print("  1. DOCX öffnen, Strg+A + F9 (Felder/TOC aktualisieren)")
print("  2. Abgabedatum auf dem Titelblatt prüfen (config.yaml: abgabe.datum)")
print(f"  3. Validator: python3 scripts/validate_docx.py {docx_path}")
