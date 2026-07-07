#!/usr/bin/env python3
"""DOCX -> HTML-Content-Preview mit IU-Format-Banner.

Macht das Dokument ohne Word/LibreOffice sichtbar: Überschriften-Hierarchie,
Fließtext mit Auszeichnungen, Tabellen, Seitenumbrüche und Sektionswechsel,
plus ein Metadaten-Banner je Sektion (Seitengröße, Ränder, Seitennummerierung).

WICHTIG: Das ist eine CONTENT-Preview. Pagination, Seitenumfang und sichtbare
Seitenzahlen entscheidet nur der echte Render (Word bzw. measure_pages.py via
LibreOffice). Die Preview ersetzt das Render-Gate aus /preflight NICHT.

Aufruf:
  python3 scripts/docx_preview.py <datei>.docx [-o <out>.html] [--open]
"""

import argparse
import html
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

CSS = """
body { background: #666; margin: 0; font-family: Arial, Helvetica, sans-serif; }
.page { background: #fff; max-width: 21cm; margin: 1em auto; padding: 2cm;
        box-shadow: 0 0 8px rgba(0,0,0,.5); }
.banner { background: #1a2a40; color: #dce6f2; font: 12px/1.5 monospace;
          max-width: 21cm; margin: 1em auto; padding: .8em 2cm; box-sizing: border-box; }
.banner b { color: #fff; }
.content p { font-size: 11pt; line-height: 1.5; margin: 0 0 6pt 0; text-align: justify; }
.content h1 { font-size: 16pt; color: #000; }
.content h2 { font-size: 14pt; color: #000; }
.content h3 { font-size: 11pt; color: #000; }
.pagebreak { border: none; border-top: 2px dashed #b00; margin: 1.5em 0; position: relative; }
.pagebreak::after { content: attr(data-label); color: #b00; font: 11px monospace;
                    position: relative; top: -1.6em; background: #fff; padding: 0 .5em; }
.tocfield { background: #fff8dc; border: 1px dashed #c90; padding: .5em;
            font: 12px monospace; color: #850; }
table.docx { border-collapse: collapse; margin: 1em 0; width: 100%; }
table.docx td, table.docx th { border: 1px solid #444; padding: 4px 8px;
                               font-size: 10pt; vertical-align: top; }
.note { max-width: 21cm; margin: 1em auto; padding: 0 2cm; box-sizing: border-box;
        color: #eee; font: 12px monospace; }
"""


def cm(emu):
    return f"{emu.cm:.2f} cm" if emu is not None else "?"


def section_banner(i, sec):
    sect_pr = sec._sectPr
    pgnum = sect_pr.find(qn("w:pgNumType"))
    fmt = (pgnum.get(qn("w:fmt")) if pgnum is not None else None) or "decimal (Default)"
    start = (pgnum.get(qn("w:start")) if pgnum is not None else None) or "fortlaufend"
    return (
        f'<div class="banner"><b>Sektion {i}</b> · '
        f"Seite: {cm(sec.page_width)} × {cm(sec.page_height)} · "
        f"Ränder: o {cm(sec.top_margin)} / u {cm(sec.bottom_margin)} / "
        f"l {cm(sec.left_margin)} / r {cm(sec.right_margin)} · "
        f"Seitenzahlen: fmt={fmt}, start={start} · "
        f"Titelseite anders: {'ja' if sec.different_first_page_header_footer else 'nein'}"
        f"</div>"
    )


def run_html(run):
    text = html.escape(run.text).replace("\n", "<br>")
    if not text:
        return ""
    if run.bold:
        text = f"<b>{text}</b>"
    if run.italic:
        text = f"<i>{text}</i>"
    if run.underline:
        text = f"<u>{text}</u>"
    return text


def is_toc_field(para):
    for instr in para._p.iter(qn("w:instrText")):
        if instr.text and "TOC" in instr.text:
            return True
    return False


def has_page_break(para):
    return any(br.get(qn("w:type")) == "page" for br in para._p.iter(qn("w:br")))


def has_section_break(para):
    ppr = para._p.find(qn("w:pPr"))
    return ppr is not None and ppr.find(qn("w:sectPr")) is not None


def para_html(para):
    parts = []
    if has_page_break(para):
        parts.append('<hr class="pagebreak" data-label="Seitenumbruch">')
    if is_toc_field(para):
        parts.append('<div class="tocfield">[TOC-Feld — Verzeichnis wird erst '
                     "in Word per Strg+A → F9 aufgebaut]</div>")
        return "".join(parts)
    style = para.style.name if para.style else ""
    inner = "".join(run_html(r) for r in para.runs) or "&nbsp;"
    if style.startswith("Heading "):
        level = min(int(style.split()[-1]), 6)
        parts.append(f"<h{level}>{inner}</h{level}>")
    else:
        parts.append(f'<p title="{html.escape(style)}">{inner}</p>')
    if has_section_break(para):
        parts.append('<hr class="pagebreak" data-label="Sektionswechsel (neue Seite)">')
    return "".join(parts)


def table_html(table):
    rows = []
    for row in table.rows:
        cells = "".join(
            "<td>" + ("".join(para_html(p) for p in cell.paragraphs) or "&nbsp;")
            + "</td>"
            for cell in row.cells
        )
        rows.append(f"<tr>{cells}</tr>")
    return '<table class="docx">' + "".join(rows) + "</table>"


def iter_block_items(doc):
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def build_html(docx_path: Path) -> str:
    doc = Document(str(docx_path))
    normal = doc.styles["Normal"].font
    head = (
        f'<div class="banner"><b>{html.escape(docx_path.name)}</b> · '
        f"Normal-Style: {normal.name or '?'} "
        f"{normal.size.pt if normal.size else '?'} pt · "
        f"{len(doc.sections)} Sektionen · {len(doc.tables)} Tabellen</div>"
    )
    banners = "".join(section_banner(i, s) for i, s in enumerate(doc.sections))
    body_parts = []
    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            body_parts.append(para_html(block))
        else:
            body_parts.append(table_html(block))
    note = (
        '<div class="note">Content-Preview (docx_preview.py) — Pagination, '
        "Seitenumfang und sichtbare Seitenzahlen sind hier NICHT verbindlich. "
        "Render-Gate: measure_pages.py bzw. Word (Strg+A → F9).</div>"
    )
    return (
        "<!DOCTYPE html><html><head><meta charset='utf-8'>"
        f"<title>Preview: {html.escape(docx_path.name)}</title>"
        f"<style>{CSS}</style></head><body>"
        f"{head}{banners}{note}"
        f'<div class="page content">{"".join(body_parts)}</div>'
        "</body></html>"
    )


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("docx", type=Path)
    ap.add_argument("-o", "--out", type=Path,
                    help="Ziel-HTML (Default: neben dem DOCX, preview_<name>.html)")
    ap.add_argument("--open", action="store_true",
                    help="Preview nach dem Erzeugen im Browser öffnen")
    args = ap.parse_args()

    if not args.docx.exists():
        sys.exit(f"FEHLER: Datei nicht gefunden: {args.docx}")
    out = args.out or args.docx.with_name(f"preview_{args.docx.stem}.html")
    out.write_text(build_html(args.docx), encoding="utf-8")
    print(f"✓ Preview: {out}")
    if args.open:
        subprocess.run(["open", str(out)], check=False)


if __name__ == "__main__":
    main()
